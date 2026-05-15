import pdfplumber
import docx
import io
import re
import spacy

nlp = spacy.load("en_core_web_sm")

class ResumeParser:
    @staticmethod
    def extract_text_and_metadata(file_bytes: bytes, filename: str) -> tuple[str, dict]:
        text = ""
        metadata = {
            "has_images": False,
            "has_tables": False,
            "pages": 0,
            "word_count": 0
        }

        if filename.endswith(".pdf"):
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                metadata["pages"] = len(pdf.pages)
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                    if len(page.images) > 0:
                        metadata["has_images"] = True
                    if len(page.find_tables()) > 0:
                        metadata["has_tables"] = True
        elif filename.endswith(".docx"):
            doc = docx.Document(io.BytesIO(file_bytes))
            metadata["pages"] = 1 # Approximation
            for para in doc.paragraphs:
                text += para.text + "\n"
            if len(doc.tables) > 0:
                metadata["has_tables"] = True
            # Checking for shapes/images in docx is complex, skipping for simplicity
        elif filename.endswith(".txt"):
            text = file_bytes.decode('utf-8')
            metadata["pages"] = 1
            
        metadata["word_count"] = len(text.split())
        return text.strip(), metadata

    @staticmethod
    def extract_entities(text: str) -> dict:
        doc = nlp(text)
        
        # Improved regex for email and phone
        email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
        # Matches formats like +1-234-567-8900, (123) 456-7890, 123.456.7890, etc.
        phone_pattern = r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
        
        emails = re.findall(email_pattern, text)
        phones = re.findall(phone_pattern, text)
        
        names = [ent.text for ent in doc.ents if ent.label_ == "PERSON"]
        
        return {
            "name": names[0] if names else None,
            "email": emails[0] if emails else None,
            "phone": phones[0] if phones else None,
        }
